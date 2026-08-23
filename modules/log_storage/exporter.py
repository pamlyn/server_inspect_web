"""巡检日志 Excel 导出工具。"""

import datetime
import io
import json
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

EXCEL_MIMETYPE = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
_HEADER_FILL = PatternFill('solid', fgColor='87CEEB')
_NORMAL_FONT = Font(name='微软雅黑')
_HEADER_FONT = Font(name='微软雅黑', bold=True)


def _excel_value(value):
    """规范化单元格值，并规避 Excel 公式注入。"""
    if value is None:
        return ''
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False, default=str)
    elif isinstance(value, (datetime.datetime, datetime.date)):
        value = value.strftime('%Y-%m-%d %H:%M:%S')
    if isinstance(value, str) and value.startswith(('=', '+', '-', '@')):
        return "'" + value
    return value


def _write_table(worksheet, columns, rows):
    """写入带统一样式、自动列宽的表格。"""
    columns = [str(column) for column in columns] or ['内容']
    for col_index, column in enumerate(columns, start=1):
        cell = worksheet.cell(row=1, column=col_index, value=column)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    for row_index, row in enumerate(rows, start=2):
        for col_index, column in enumerate(columns, start=1):
            if isinstance(row, (list, tuple)):
                value = row[col_index - 1] if col_index <= len(row) else ''
            elif isinstance(row, dict):
                value = row.get(column, '')
            else:
                value = row if col_index == 1 else ''
            cell = worksheet.cell(row=row_index, column=col_index, value=_excel_value(value))
            cell.font = _NORMAL_FONT
            cell.alignment = Alignment(vertical='top', wrap_text=True)

    for col_index, column in enumerate(columns, start=1):
        max_length = len(column)
        for row_index in range(2, worksheet.max_row + 1):
            value = str(worksheet.cell(row=row_index, column=col_index).value or '')
            max_length = max(max_length, max(len(line) for line in (value.splitlines() or [''])))
        worksheet.column_dimensions[get_column_letter(col_index)].width = min(max(max_length + 2, 10), 50)
    worksheet.freeze_panes = 'A2'


def _save_workbook(workbook, filename_prefix):
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    safe_name = re.sub(r'[\\/:*?"<>|]+', '_', str(filename_prefix)).strip() or '巡检日志'
    filename = f'{safe_name}_{datetime.datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
    return output, filename


def build_log_list_excel(logs):
    """生成巡检日志列表工作簿。"""
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = '巡检日志'
    columns = ['时间', '类型', '来源', '状态', '巡检对象', '摘要', '操作人', '记录数', '耗时']
    rows = []
    for log in logs:
        rows.append([
            log.get('start_time') or log.get('created_at') or '',
            log.get('type_label') or log.get('inspection_type') or '',
            log.get('source_label') or log.get('trigger_source') or '',
            log.get('status_label') or log.get('status') or '',
            log.get('target_label') or log.get('target') or '',
            log.get('summary') or '', log.get('operator') or '系统',
            log.get('record_count') if log.get('record_count') is not None else '',
            log.get('duration') or '',
        ])
    _write_table(worksheet, columns, rows)
    return _save_workbook(workbook, '巡检日志')
