"""
自定义脚本模块 - 路由定义
"""

from flask import Blueprint, jsonify, request, send_file, session
import datetime
import io
import re
import time
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from modules.auth.helpers import ADMIN_USERNAME, can_access_custom_script, login_required
from modules.custom_scripts.helpers import custom_scripts, script_id_counter, save_scripts, format_sql_value, resolve_period_value, get_variable_value, allocate_script_id
from modules.custom_scripts.cross_db import compare_cross_db, cross_db_summary_text, cross_db_config_to_text
from modules.inspection.helpers import execute_sql
from modules.config_mgmt.helpers import get_config

custom_scripts_bp = Blueprint('custom_scripts', __name__, url_prefix='/api/custom_scripts')


def _can_manage_scripts():
    return session.get('username') == ADMIN_USERNAME


def _can_access_script(script):
    return can_access_custom_script(session.get('username', ''), script)


def _forbidden(message='无权访问该自定义SQL'):
    return jsonify({'error': message}), 403


@custom_scripts_bp.route('', methods=['GET', 'POST'])
@login_required
def custom_scripts_api():
    """自定义SQL脚本管理"""
    global custom_scripts, script_id_counter

    if request.method == 'GET':
        scripts = [dict(script, category=(script.get('category') or '').strip() or '未分类')
                   for script in custom_scripts if _can_access_script(script)]
        return jsonify({'scripts': scripts})
    if not _can_manage_scripts():
        return _forbidden('仅管理员可新增或编辑自定义SQL脚本')
    elif request.method == 'POST':
        data = request.json
        script_id = data.get('id')
        name = data.get('name')
        mode = data.get('mode', 'single_db')
        variables = data.get('variables', [])
        rules = data.get('rules', [])
        cross_db_config = data.get('cross_db_config', {})
        category = (data.get('category') or '').strip() or '未分类'

        if not name:
            return jsonify({'error': '脚本名称不能为空'}), 400

        if script_id:
            script = next((s for s in custom_scripts if s['id'] == script_id), None)
            if not script:
                return jsonify({'error': '脚本不存在'}), 404

            script['name'] = name
            script['category'] = category
            script['mode'] = mode
            script['variables'] = variables
            script['rules'] = rules
            script['scheduled'] = data.get('scheduled', False)
            script['daily'] = data.get('daily', False)
            script['realtime'] = data.get('realtime', False)
            script['updated_at'] = datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')

            if mode == 'single_db':
                script['database'] = data.get('database')
                script['content'] = data.get('content')
                script.pop('source_db', None)
                script.pop('target_db', None)
                script.pop('source_sql', None)
                script.pop('target_sql', None)
                script.pop('cross_db_config', None)
            else:
                script['source_db'] = data.get('source_db')
                script['target_db'] = data.get('target_db')
                script['source_sql'] = data.get('source_sql')
                script['target_sql'] = data.get('target_sql')
                script['cross_db_config'] = cross_db_config
                script.pop('database', None)
                script.pop('content', None)
        else:
            new_script = {
                'id': allocate_script_id(),
                'name': name,
                'category': category,
                'mode': mode,
                'scheduled': data.get('scheduled', False),
                'daily': data.get('daily', False),
                'realtime': data.get('realtime', False),
                'variables': variables,
                'rules': rules,
                'created_at': datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')
            }

            if mode == 'single_db':
                new_script['database'] = data.get('database')
                new_script['content'] = data.get('content')
                if not new_script['content']:
                    return jsonify({'error': '脚本内容不能为空'}), 400
            else:
                new_script['source_db'] = data.get('source_db')
                new_script['target_db'] = data.get('target_db')
                new_script['source_sql'] = data.get('source_sql')
                new_script['target_sql'] = data.get('target_sql')
                new_script['cross_db_config'] = cross_db_config
                if not new_script['source_sql'] or not new_script['target_sql']:
                    return jsonify({'error': '源数据库和目标数据库SQL不能为空'}), 400

            custom_scripts.append(new_script)

        save_scripts()
        return jsonify({'message': '脚本保存成功'})


@custom_scripts_bp.route('/<script_id>', methods=['GET', 'DELETE', 'PUT'])
@login_required
def custom_script_api(script_id):
    """单个自定义脚本管理"""
    script = next((s for s in custom_scripts if s['id'] == script_id), None)
    if not script:
        return jsonify({'error': '脚本不存在'}), 404
    if not _can_access_script(script):
        return _forbidden()

    if request.method == 'GET':
        response_script = dict(script, category=(script.get('category') or '').strip() or '未分类')
        return jsonify({'script': response_script})
    if not _can_manage_scripts():
        return _forbidden('仅管理员可编辑或删除自定义SQL脚本')
    elif request.method == 'DELETE':
        # 必须就地删除。`from helpers import custom_scripts` 拿到的是同一个列表对象，
        # 重新赋值只会改本模块的名字，helpers 和看板模块仍指向旧列表，
        # 之后新增的脚本在看板里就选不到了。
        custom_scripts[:] = [s for s in custom_scripts if s['id'] != script_id]
        save_scripts()
        return jsonify({'message': '脚本删除成功'})
    elif request.method == 'PUT':
        data = request.json
        script['name'] = data.get('name', script['name'])
        script['category'] = (data.get('category') or script.get('category') or '未分类').strip() or '未分类'
        script['database'] = data.get('database', script['database'])
        script['content'] = data.get('content', script['content'])
        script['variables'] = data.get('variables', script.get('variables', []))
        script['rules'] = data.get('rules', script.get('rules', []))
        script['scheduled'] = data.get('scheduled', script.get('scheduled', False))
        script['daily'] = data.get('daily', script.get('daily', False))
        script['realtime'] = data.get('realtime', script.get('realtime', False))
        script['updated_at'] = datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        save_scripts()
        return jsonify({'message': '脚本更新成功'})


@custom_scripts_bp.route('/export', methods=['POST'])
@login_required
def export_custom_script_result():
    """导出自定义SQL结果为列宽自适应的Excel文件。"""
    data = request.get_json() or {}
    columns = data.get('columns') or []
    rows = data.get('rows') or []
    script_name = data.get('script_name') or 'SQL结果'
    header_color = str(data.get('header_color') or '#87CEEB').lstrip('#').upper()
    if not re.fullmatch(r'[0-9A-F]{6}', header_color):
        header_color = '87CEEB'
    if not columns:
        return jsonify({'error': '没有可导出的列'}), 400

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = 'SQL结果'
    header_fill = PatternFill('solid', fgColor=header_color)
    normal_font = Font(name='微软雅黑')
    header_font = Font(name='微软雅黑', bold=True)
    for col_index, column in enumerate(columns, start=1):
        cell = worksheet.cell(row=1, column=col_index, value=str(column))
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    for row_index, row in enumerate(rows, start=2):
        for col_index, column in enumerate(columns, start=1):
            value = row[col_index - 1] if isinstance(row, list) and col_index <= len(row) else row.get(column) if isinstance(row, dict) else ''
            if value is None:
                value = ''
            elif isinstance(value, str) and value.startswith(('=', '+', '-', '@')):
                value = "'" + value
            cell = worksheet.cell(row=row_index, column=col_index, value=value)
            cell.font = normal_font
            cell.alignment = Alignment(vertical='top', wrap_text=True)

    for col_index, column in enumerate(columns, start=1):
        max_length = len(str(column))
        for row_index in range(2, worksheet.max_row + 1):
            max_length = max(max_length, max(len(line) for line in str(worksheet.cell(row_index, col_index).value or '').splitlines()))
        worksheet.column_dimensions[get_column_letter(col_index)].width = min(max(max_length + 2, 10), 50)
    worksheet.freeze_panes = 'A2'

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    safe_name = re.sub(r'[\\/:*?"<>|]+', '_', str(script_name)).strip() or 'SQL结果'
    filename = f'{safe_name}_{datetime.datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
    return send_file(output, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True, download_name=filename)


@custom_scripts_bp.route('/<script_id>/execute', methods=['POST'])
@login_required
def execute_custom_script(script_id):
    """执行已保存的自定义脚本"""
    global custom_scripts
    from flask import session

    start_time = datetime.datetime.now()
    start_ts = time.perf_counter()
    operator = session.get('username')

    script = next((s for s in custom_scripts if s['id'] == script_id), None)
    if not script:
        return jsonify({'error': '脚本不存在'}), 404
    if not _can_access_script(script):
        return _forbidden()

    script_name = script.get('name', script_id)
    mode = script.get('mode', 'single_db')

    try:
        params = request.get_json() or {}
        DATABASE_CONFIG = get_config('databaseConfig')

        if mode == 'single_db':
            sql_content = script.get('content', '')
            variables = script.get('variables', [])

            for var in variables:
                var_name = var.get('name', '')
                var_type = var.get('type', 'text')
                var_value = get_variable_value(var, params)
                formatted_value = format_sql_value(var, var_value)
                sql_content = sql_content.replace(f'#{{{var_name}}}', formatted_value)

            database = script.get('database')
            db_config = DATABASE_CONFIG.get(database)
            if not db_config:
                _log_custom_script(start_time, start_ts, operator, script_name, mode, 0, None, '数据库配置不存在')
                return jsonify({'error': '数据库配置不存在'}), 400

            columns, rows = execute_sql(db_config['type'], db_config, sql_content)
            if columns is None:
                _log_custom_script(start_time, start_ts, operator, script_name, mode, 0, None, rows)
                return jsonify({'error': rows}), 500

            record_count = len(rows) if rows else 0
            _log_custom_script(start_time, start_ts, operator, script_name, mode, record_count,
                               {'columns': columns, 'rows': rows}, None)
            return jsonify({'columns': columns, 'rows': rows})
        else:
            source_db = script.get('source_db')
            target_db = script.get('target_db')
            source_sql = script.get('source_sql', '')
            target_sql = script.get('target_sql', '')
            variables = script.get('variables', [])

            for var in variables:
                var_name = var.get('name', '')
                var_type = var.get('type', 'text')
                var_value = get_variable_value(var, params)
                formatted_value = format_sql_value(var, var_value)
                source_sql = source_sql.replace(f'#{{{var_name}}}', formatted_value)
                target_sql = target_sql.replace(f'#{{{var_name}}}', formatted_value)

            source_db_config = DATABASE_CONFIG.get(source_db)
            target_db_config = DATABASE_CONFIG.get(target_db)

            if not source_db_config or not target_db_config:
                _log_custom_script(start_time, start_ts, operator, script_name, mode, 0, None, '数据库配置不存在')
                return jsonify({'error': '数据库配置不存在'}), 400

            source_columns, source_rows = execute_sql(source_db_config['type'], source_db_config, source_sql)
            if source_columns is None:
                _log_custom_script(start_time, start_ts, operator, script_name, mode, 0, None, f'源数据库查询失败: {source_rows}')
                return jsonify({'error': f'源数据库查询失败: {source_rows}'}), 500

            target_columns, target_rows = execute_sql(target_db_config['type'], target_db_config, target_sql)
            if target_columns is None:
                _log_custom_script(start_time, start_ts, operator, script_name, mode, 0, None, f'目标数据库查询失败: {target_rows}')
                return jsonify({'error': f'目标数据库查询失败: {target_rows}'}), 500

            cross_db_config = script.get('cross_db_config', {}) or {}
            compare_result = compare_cross_db(source_columns, source_rows, target_columns, target_rows, cross_db_config)
            summary = compare_result.get('summary', {})
            is_consistent = compare_result.get('is_consistent', True)

            record_count = len(compare_result.get('rows', []))
            result_for_log = {
                'columns': compare_result.get('columns'),
                'rows': compare_result.get('rows'),
                'summary': summary,
                'is_consistent': is_consistent,
                'cross_db_config': cross_db_config,
                'source_db': source_db,
                'target_db': target_db,
                'script_name': script_name,
                'mode': 'cross_db',
            }
            status = 'success' if is_consistent else 'warning'
            _log_custom_script(start_time, start_ts, operator, script_name, mode, record_count,
                               result_for_log, None, summary=summary, status=status,
                               extra_summary=cross_db_summary_text(summary))
            return jsonify({
                'columns': compare_result.get('columns'),
                'rows': compare_result.get('rows'),
                'summary': summary,
                'is_consistent': is_consistent,
            })
    except Exception as e:
        print(f"执行脚本失败: {e}")
        import traceback
        traceback.print_exc()
        _log_custom_script(start_time, start_ts, operator, script_name, mode, 0, None, str(e))
        return jsonify({'error': str(e)}), 500


def _log_custom_script(start_time, start_ts, operator, script_name, mode, record_count, result, error,
                       summary=None, status=None, extra_summary=''):
    """记录自定义脚本执行日志（成功/失败统一入口）。

    :param summary: 可选，跨库对比的 summary dict（用于日志 result）
    :param status:  可选，覆盖默认状态（如跨库对比不一致时为 warning）
    :param extra_summary: 可选，附加到摘要文本的补充说明（如跨库对比中文统计）
    """
    from modules.log_storage.helpers import record_inspection_log
    end_time = datetime.datetime.now()
    if status is None:
        status = 'error' if error else 'success'
    mode_label = '跨库对比' if mode == 'cross_db' else '单库'
    if error:
        summary_text = f"自定义脚本失败[{mode_label}]: {script_name}"
    else:
        base = f"自定义脚本完成[{mode_label}]: {script_name}，{record_count} 条记录"
        summary_text = f"{base}（{extra_summary}）" if extra_summary else base

    # 跨库对比把 summary 统计并入 result，便于日志页解析
    log_result = result
    if summary and isinstance(log_result, dict):
        log_result = dict(log_result)
        log_result['summary'] = summary

    record_inspection_log(
        inspection_type='custom_script', trigger_source='manual',
        target=script_name, operator=operator, status=status,
        start_time=start_time, end_time=end_time,
        duration=time.perf_counter() - start_ts,
        summary=summary_text, record_count=record_count,
        result=log_result, error=error,
    )


# Note: test_custom_script路由注册在app_new.py中，因为它需要在
# /api/test_custom_script路径下，而不是在Blueprint前缀/api/custom_scripts下